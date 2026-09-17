from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "assets" / "build_minutes.py"
VERIFIER = ROOT / "scripts" / "verify_docx_format.py"
REFERENCE = ROOT / "assets" / "template_minutes.docx"
FILENAME = "测试项目_某机构专家访谈纪要_20260911.docx"
PAYLOAD = {
    "project": "测试项目",
    "subject": "某机构",
    "interview_type": "专家访谈",
    "date_compact": "20260911",
    "meta": {
        "访谈时间": "2026年9月11日",
        "访谈地点": "线上",
        "访谈来源": "公司",
        "访谈对象及其背景介绍": "某行业从业人员",
        "访谈人员": "张三、李四",
    },
    "conclusions": [{"lead": "核心结论：", "body": "结论正文。"}],
    "sections": [
        {"kind": "section", "text": "行业情况"},
        {"kind": "subsection", "text": "市场概况"},
        {"kind": "point", "lead": "市场特征：", "body": "正文内容。"},
    ],
}


def run_verifier(path: Path):
    return subprocess.run(
        [
            sys.executable,
            str(VERIFIER),
            str(path),
            "--reference",
            str(REFERENCE),
            "--project",
            "测试项目",
            "--subject",
            "某机构",
            "--interview-type",
            "专家访谈",
            "--date-compact",
            "20260911",
            "--expected-meta",
            json.dumps(PAYLOAD["meta"], ensure_ascii=False),
        ],
        text=True,
        capture_output=True,
    )


def copy_valid(source: Path, target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / FILENAME
    shutil.copy2(source, target)
    return target


class FormatGuardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        payload_path = cls.root / "payload.json"
        payload_path.write_text(json.dumps(PAYLOAD, ensure_ascii=False), encoding="utf-8")
        output_dir = cls.root / "baseline"
        result = subprocess.run(
            [sys.executable, str(BUILDER), "--payload", str(payload_path), "--output-dir", str(output_dir)],
            text=True,
            capture_output=True,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stdout + result.stderr)
        cls.baseline = output_dir / FILENAME

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_valid_output_passes(self):
        result = run_verifier(self.baseline)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("FORMAT CHECK PASSED", result.stdout)

    def test_wrong_filename_is_blocked(self):
        candidate = self.root / "wrong-name.docx"
        shutil.copy2(self.baseline, candidate)
        result = run_verifier(candidate)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("文件名错误", result.stderr)

    def test_global_font_override_is_blocked(self):
        candidate = copy_valid(self.baseline, self.root / "font-corrupt")
        temporary = candidate.with_suffix(".tmp")
        with zipfile.ZipFile(candidate) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/styles.xml":
                    data = data.replace("楷体".encode(), "Hiragino Sans GB".encode())
                target.writestr(info, data)
        temporary.replace(candidate)
        result = run_verifier(candidate)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("word/styles.xml", result.stderr)
        self.assertIn("Hiragino Sans GB", result.stderr)

    def test_manual_numbering_is_blocked(self):
        candidate = copy_valid(self.baseline, self.root / "number-corrupt")
        doc = Document(candidate)
        section = next(
            paragraph
            for paragraph in doc.paragraphs
            if (node := paragraph._p.find(".//" + qn("w:numId"))) is not None
            and node.get(qn("w:val")) == "13"
        )
        section.runs[0].text = "一 " + section.runs[0].text
        doc.save(candidate)
        result = run_verifier(candidate)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("手写编号", result.stderr)

    def test_placeholder_metadata_is_blocked(self):
        candidate = copy_valid(self.baseline, self.root / "meta-corrupt")
        doc = Document(candidate)
        paragraph = next(p for p in doc.tables[0].cell(0, 0).paragraphs if p.text.startswith("访谈人员"))
        paragraph.runs[-1].text = "111"
        doc.save(candidate)
        result = run_verifier(candidate)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("占位内容", result.stderr)

    def test_custom_reference_preserves_critical_parts(self):
        custom_reference = self.root / "custom-reference.docx"
        temporary = custom_reference.with_suffix(".tmp")
        with zipfile.ZipFile(REFERENCE) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/styles.xml":
                    data = data.replace(b"</w:styles>", b"\n</w:styles>")
                target.writestr(info, data)
        temporary.replace(custom_reference)
        payload_path = self.root / "custom-payload.json"
        payload_path.write_text(json.dumps(PAYLOAD, ensure_ascii=False), encoding="utf-8")
        output_dir = self.root / "custom-output"
        result = subprocess.run(
            [
                sys.executable,
                str(BUILDER),
                "--payload",
                str(payload_path),
                "--reference",
                str(custom_reference),
                "--output-dir",
                str(output_dir),
            ],
            text=True,
            capture_output=True,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        with zipfile.ZipFile(custom_reference) as reference_zip, zipfile.ZipFile(output_dir / FILENAME) as output_zip:
            for part in ("word/styles.xml", "word/numbering.xml", "word/settings.xml", "word/theme/theme1.xml"):
                self.assertEqual(reference_zip.read(part), output_zip.read(part), part)


if __name__ == "__main__":
    unittest.main()
