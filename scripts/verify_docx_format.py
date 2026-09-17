#!/usr/bin/env python3
"""Fail-closed format verifier for Interview Minutes DOCX outputs."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = ROOT / "assets" / "template_minutes.docx"
DEFAULT_MANIFEST = ROOT / "assets" / "template_manifest.json"
CRITICAL_PARTS = (
    "word/styles.xml",
    "word/numbering.xml",
    "word/settings.xml",
    "word/theme/theme1.xml",
)
META_LABELS = (
    "访谈时间",
    "访谈地点",
    "访谈来源",
    "访谈对象及其背景介绍",
    "访谈人员",
)
PLACEHOLDER_RE = re.compile(
    r"(?:^|\s)(?:111|123|xxx+|tbd|todo|待补充|待填写|访谈团队|yyyy年m月d日)(?:\s|$)",
    re.IGNORECASE,
)
MANUAL_PREFIX = {
    "13": re.compile(r"^\s*(?:第?[一二三四五六七八九十]+[、.．\s]|\d+[、.．)）\s])"),
    "10": re.compile(r"^\s*(?:\d+[、.．)）\s]|[（(]\d+[）)])"),
    "7": re.compile(r"^\s*(?:\d+[、.．)）\s]|[（(]\d+[）)])"),
    "4": re.compile(r"^\s*(?:[-–—]\s+|[•➢]|\d+[、.．)）\s]|[（(]\d+[）)])"),
}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def zip_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def canonical_xml(element) -> bytes:
    if element is None:
        return b""
    node = deepcopy(element)
    for descendant in node.iter():
        for attr in list(descendant.attrib):
            if etree.QName(attr).localname.startswith("rsid"):
                del descendant.attrib[attr]
    return etree.tostring(node, method="c14n")


def paragraph_properties_match(candidate, reference, *, ignore_page_break=False) -> bool:
    candidate_node = deepcopy(candidate)
    reference_node = deepcopy(reference)
    if ignore_page_break:
        for node in (candidate_node, reference_node):
            if node is not None:
                for page_break in list(node.findall(qn("w:pageBreakBefore"))):
                    node.remove(page_break)
    return canonical_xml(candidate_node) == canonical_xml(reference_node)


def num_id(paragraph) -> str | None:
    node = paragraph._p.find(".//" + qn("w:numId"))
    return node.get(qn("w:val")) if node is not None else None


def paragraph_role(paragraph, *, is_first: bool = False) -> str | None:
    if is_first:
        return "title"
    if paragraph.style and paragraph.style.style_id == "1":
        return "minutes_h1" if paragraph.text.strip() == "访谈纪要" else "main_h1"
    nid = num_id(paragraph)
    return {"7": "conclusion", "13": "section", "10": "subsection", "4": "point"}.get(nid)


def reference_prototypes(doc: Document):
    paragraphs = doc.paragraphs
    return {
        "title": paragraphs[0],
        "main_h1": next(p for p in paragraphs if p.text.strip() == "主要结论"),
        "minutes_h1": next(p for p in paragraphs if p.text.strip() == "访谈纪要"),
        "conclusion": next(p for p in paragraphs if num_id(p) == "7"),
        "section": next(p for p in paragraphs if num_id(p) == "13"),
        "subsection": next(p for p in paragraphs if num_id(p) == "10"),
        "point": next(p for p in paragraphs if num_id(p) == "4"),
    }


def run_property_set(paragraph) -> set[bytes]:
    return {canonical_xml(run._r.rPr) for run in paragraph.runs}


def all_declared_fonts(parts: dict[str, bytes]) -> set[str]:
    fonts: set[str] = set()
    pattern = re.compile(rb'w:(?:ascii|hAnsi|eastAsia|cs)="([^"]+)"')
    for name in ("word/document.xml", "word/styles.xml"):
        for raw in pattern.findall(parts.get(name, b"")):
            fonts.add(raw.decode("utf-8", errors="replace"))
    return fonts


def parse_meta(doc: Document) -> dict[str, str]:
    result: dict[str, str] = {}
    if not doc.tables:
        return result
    for paragraph in doc.tables[0].cell(0, 0).paragraphs:
        text = paragraph.text.strip()
        for label in META_LABELS:
            match = re.match(rf"^{re.escape(label)}[：:]\s*(.*)$", text)
            if match:
                result[label] = match.group(1).strip()
                break
    return result


def expected_filename(project: str, subject: str, interview_type: str, date_compact: str) -> str:
    return f"{project}_{subject}{interview_type}纪要_{date_compact}.docx"


def validate(
    candidate: Path,
    reference: Path,
    *,
    project: str | None = None,
    subject: str | None = None,
    interview_type: str = "专家访谈",
    date_compact: str | None = None,
    expected_meta: dict[str, str] | None = None,
) -> list[str]:
    errors: list[str] = []
    if not candidate.is_file():
        return [f"候选文件不存在：{candidate}"]
    if not reference.is_file():
        return [f"参考模板不存在：{reference}"]

    candidate_parts = zip_parts(candidate)
    reference_parts = zip_parts(reference)

    if reference.resolve() == DEFAULT_REFERENCE.resolve() and DEFAULT_MANIFEST.is_file():
        manifest = json.loads(DEFAULT_MANIFEST.read_text(encoding="utf-8"))
        for part, expected_hash in manifest.get("critical_parts", {}).items():
            actual_hash = sha256(reference_parts.get(part, b""))
            if actual_hash != expected_hash:
                errors.append(f"默认企业模板指纹异常：{part}")

    for part in CRITICAL_PARTS:
        if candidate_parts.get(part) != reference_parts.get(part):
            errors.append(f"关键模板部件被改写：{part}")

    candidate_fonts = all_declared_fonts(candidate_parts)
    allowed_fonts = all_declared_fonts(reference_parts)
    unexpected_fonts = sorted(candidate_fonts - allowed_fonts)
    if unexpected_fonts:
        errors.append("出现模板外字体：" + "、".join(unexpected_fonts))

    candidate_doc = Document(candidate)
    reference_doc = Document(reference)
    prototypes = reference_prototypes(reference_doc)

    if len(candidate_doc.tables) != len(reference_doc.tables):
        errors.append("元信息表数量与模板不一致")
    elif candidate_doc.tables:
        if canonical_xml(candidate_doc.tables[0]._tbl.tblPr) != canonical_xml(reference_doc.tables[0]._tbl.tblPr):
            errors.append("元信息表格属性被改写")
        candidate_cell = candidate_doc.tables[0].cell(0, 0)
        reference_cell = reference_doc.tables[0].cell(0, 0)
        if canonical_xml(candidate_cell._tc.tcPr) != canonical_xml(reference_cell._tc.tcPr):
            errors.append("元信息单元格属性被改写")

        reference_meta = reference_cell.paragraphs
        reference_meta_ppr = {canonical_xml(p._p.pPr) for p in reference_meta}
        reference_meta_rpr = {canonical_xml(r._r.rPr) for p in reference_meta for r in p.runs}
        for paragraph in candidate_cell.paragraphs:
            if canonical_xml(paragraph._p.pPr) not in reference_meta_ppr:
                errors.append(f"元信息段落格式被改写：{paragraph.text[:24]}")
            for run in paragraph.runs:
                if canonical_xml(run._r.rPr) not in reference_meta_rpr:
                    errors.append(f"元信息字体或粗体属性被改写：{paragraph.text[:24]}")
                    break

    if candidate_doc.sections and reference_doc.sections:
        if canonical_xml(candidate_doc.sections[0]._sectPr) != canonical_xml(reference_doc.sections[0]._sectPr):
            errors.append("页面尺寸、页边距或节属性与模板不一致")

    nonempty_roles: list[str] = []
    for index, paragraph in enumerate(candidate_doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        role = paragraph_role(paragraph, is_first=(index == 0))
        if role is None:
            errors.append(f"发现未继承模板层级的段落：{text[:30]}")
            continue
        nonempty_roles.append(role)
        prototype = prototypes[role]
        if not paragraph_properties_match(
            paragraph._p.pPr,
            prototype._p.pPr,
            ignore_page_break=(role == "minutes_h1"),
        ):
            errors.append(f"段落属性与模板不一致：{text[:30]}")
        allowed_rpr = run_property_set(prototype)
        for run in paragraph.runs:
            if canonical_xml(run._r.rPr) not in allowed_rpr:
                errors.append(f"字体、字号或粗体范围与模板不一致：{text[:30]}")
                break
        nid = num_id(paragraph)
        if nid in MANUAL_PREFIX and MANUAL_PREFIX[nid].search(text):
            errors.append(f"自动编号段落含手写编号或项目符号：{text[:30]}")

    for required_role in ("title", "main_h1", "minutes_h1", "conclusion", "section", "subsection", "point"):
        if required_role not in nonempty_roles:
            errors.append(f"缺少必要结构：{required_role}")

    if candidate_doc.paragraphs:
        title = candidate_doc.paragraphs[0].text.strip()
        if project and subject:
            expected_title = f"{project}-{subject}{interview_type}"
            if title != expected_title:
                errors.append(f"标题错误：应为“{expected_title}”，实际为“{title}”")
        if PLACEHOLDER_RE.search(title.lower()):
            errors.append("标题仍含占位内容")

    if project and subject and date_compact:
        expected_name = expected_filename(project, subject, interview_type, date_compact)
        if candidate.name != expected_name:
            errors.append(f"文件名错误：应为“{expected_name}”，实际为“{candidate.name}”")

    meta = parse_meta(candidate_doc)
    for label in META_LABELS:
        if label not in meta:
            errors.append(f"元信息缺少字段：{label}")
    for label, value in meta.items():
        if PLACEHOLDER_RE.search(value.lower()):
            errors.append(f"元信息含占位内容：{label}={value}")
    if expected_meta:
        for label, expected_value in expected_meta.items():
            actual_value = meta.get(label, "")
            if actual_value != expected_value:
                errors.append(f"元信息不一致：{label}应为“{expected_value}”，实际为“{actual_value}”")

    text_all = "\n".join(p.text for p in candidate_doc.paragraphs)
    if "访谈纪要" in text_all:
        minutes = next((p for p in candidate_doc.paragraphs if p.text.strip() == "访谈纪要"), None)
        if minutes is not None and minutes._p.pPr is not None:
            if minutes._p.pPr.find(qn("w:pageBreakBefore")) is None:
                errors.append("“访谈纪要”未设置模板规定的受控分页")

    return list(dict.fromkeys(errors))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("candidate", type=Path)
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--project")
    parser.add_argument("--subject")
    parser.add_argument("--interview-type", default="专家访谈")
    parser.add_argument("--date-compact")
    parser.add_argument("--expected-meta", help="JSON object keyed by metadata label without colon")
    parser.add_argument("--report-json", type=Path)
    args = parser.parse_args()

    expected_meta = json.loads(args.expected_meta) if args.expected_meta else None
    errors = validate(
        args.candidate,
        args.reference,
        project=args.project,
        subject=args.subject,
        interview_type=args.interview_type,
        date_compact=args.date_compact,
        expected_meta=expected_meta,
    )
    report = {"candidate": str(args.candidate), "reference": str(args.reference), "passed": not errors, "errors": errors}
    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if errors:
        print("FORMAT CHECK FAILED", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    print("FORMAT CHECK PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
