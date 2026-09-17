# -*- coding: utf-8 -*-
"""机构级模板访谈纪要生成脚手架。

本脚手架只替换内容，字体、字号、缩进、间距、编号、粗体范围和页面结构均
从参考文档原生 OOXML 深拷贝，不通过 set_run_font/kai 等函数重设格式。
"""

from copy import deepcopy
import os
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
TEMPLATE = Path(os.environ.get('REFERENCE', HERE / 'template_minutes.docx'))
OUT = Path(os.environ.get('OUT', HERE / '访谈纪要_输出.docx'))

TITLE = '项目-访谈对象-专家访谈'
META = [
    ('访谈时间：', 'YYYY年M月D日'),
    ('访谈地点：', '线上'),
    ('访谈来源：', ''),
    ('访谈对象及其背景介绍：', '姓名、职务及与本次主题相关的背景'),
    ('访谈人员：', '访谈团队'),
]
CONCLUSIONS = [
    ('结论引导语：', '核心结论及其直接依据。'),
]
SECTIONS = [
    ('section', '一级章节标题'),
    ('subsection', '二级标题'),
    ('point', ('要点引导语：', '要点正文。')),
]


def number_id(paragraph):
    node = paragraph._p.find('.//' + qn('w:numId'))
    return node.get(qn('w:val')) if node is not None else None


def clear_inline_content(paragraph_element):
    for child in list(paragraph_element):
        if child.tag != qn('w:pPr'):
            paragraph_element.remove(child)


def clone_run(run_element, text):
    new_run = deepcopy(run_element)
    for child in list(new_run):
        if child.tag != qn('w:rPr'):
            new_run.remove(child)
    text_node = OxmlElement('w:t')
    if text.startswith(' ') or text.endswith(' '):
        text_node.set(qn('xml:space'), 'preserve')
    text_node.text = text
    new_run.append(text_node)
    return new_run


def clone_simple(paragraph, text, run_index=0):
    new_p = deepcopy(paragraph._p)
    run_template = paragraph.runs[run_index]._r
    clear_inline_content(new_p)
    new_p.append(clone_run(run_template, text))
    return new_p


def clone_split(paragraph, lead, body):
    new_p = deepcopy(paragraph._p)
    clear_inline_content(new_p)
    new_p.append(clone_run(paragraph.runs[0]._r, lead))
    if body:
        body_template = paragraph.runs[1]._r if len(paragraph.runs) > 1 else paragraph.runs[0]._r
        new_p.append(clone_run(body_template, body))
    return new_p


def add_page_break_before(paragraph_element):
    ppr = paragraph_element.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        paragraph_element.insert(0, ppr)
    if ppr.find(qn('w:pageBreakBefore')) is None:
        page_break = OxmlElement('w:pageBreakBefore')
        insertion_index = 0
        allowed_before = {qn('w:pStyle'), qn('w:keepNext'), qn('w:keepLines')}
        for index, child in enumerate(list(ppr)):
            if child.tag in allowed_before:
                insertion_index = index + 1
            else:
                break
        ppr.insert(insertion_index, page_break)


def clone_meta(paragraph, label, value, value_run=None):
    new_p = deepcopy(paragraph._p)
    clear_inline_content(new_p)
    new_p.append(clone_run(paragraph.runs[0]._r, label))
    if value:
        candidate = value_run or next((r for r in paragraph.runs[1:] if not r.bold), paragraph.runs[-1])
        new_p.append(clone_run(candidate._r, value))
    return new_p


def find_meta_prototype(paragraphs, label):
    stem = label.rstrip('：:')
    return next((p for p in paragraphs if p.text.startswith(stem)), None)


def build():
    doc = Document(TEMPLATE)
    paragraphs = doc.paragraphs
    body = doc.element.body
    table_element = next(child for child in body if child.tag == qn('w:tbl'))
    section_properties = body.find(qn('w:sectPr'))

    prototype = {
        'title': paragraphs[0],
        'h1': next(p for p in paragraphs if p.style.name == 'Heading 1'),
        'empty': next((p for p in paragraphs if not p.text and number_id(p) is None), None),
        'conclusion': next(p for p in paragraphs if number_id(p) == '7'),
        'section': next(p for p in paragraphs if number_id(p) == '13'),
        'subsection': next(p for p in paragraphs if number_id(p) == '10'),
        'point': next(p for p in paragraphs if number_id(p) == '4'),
    }

    for child in list(body):
        if child.tag == qn('w:p'):
            body.remove(child)
    body.insert(body.index(table_element), clone_simple(prototype['title'], TITLE))

    def append(element):
        if section_properties is not None:
            section_properties.addprevious(element)
        else:
            body.append(element)

    append(clone_simple(prototype['h1'], '主要结论'))
    for lead, text in CONCLUSIONS:
        append(clone_split(prototype['conclusion'], lead, text))

    if prototype['empty'] is not None:
        append(deepcopy(prototype['empty']._p))

    minutes_heading = clone_simple(prototype['h1'], '访谈纪要')
    add_page_break_before(minutes_heading)
    append(minutes_heading)

    for kind, payload in SECTIONS:
        if kind in {'section', 'subsection'}:
            append(clone_simple(prototype[kind], payload))
        elif kind == 'point':
            lead, text = payload
            append(clone_split(prototype['point'], lead, text))
        elif kind == 'empty':
            if prototype['empty'] is None:
                append(OxmlElement('w:p'))
            else:
                append(deepcopy(prototype['empty']._p))
        else:
            raise ValueError(f'未知段落类型：{kind}')

    cell = doc.tables[0].rows[0].cells[0]
    old_meta = list(cell.paragraphs)
    fallback_value = next(
        (r for p in old_meta for r in p.runs if r.text.strip() and not r.bold),
        old_meta[0].runs[0],
    )
    new_meta = []
    for label, value in META:
        paragraph = find_meta_prototype(old_meta, label) or old_meta[0]
        value_template = next((r for r in paragraph.runs[1:] if not r.bold), fallback_value)
        new_meta.append(clone_meta(paragraph, label, value, value_template))
    for node in list(cell._tc.findall(qn('w:p'))):
        cell._tc.remove(node)
    for node in new_meta:
        cell._tc.append(node)

    properties = doc.core_properties
    properties.author = 'Interview Minutes contributors'
    properties.last_modified_by = 'Interview Minutes contributors'
    properties.title = TITLE
    properties.subject = 'Institutional interview minutes'
    properties.keywords = 'interview, minutes, due diligence'
    properties.comments = ''

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f'SAVED {OUT}')


if __name__ == '__main__':
    build()
